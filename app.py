#!/usr/bin/env python3
"""Golden Set 검수 Viewer 백엔드 (API 계약: API.md).

  python3 app.py [--e2e ../e2e] [--pre-dir out/ao-pre-image] [--host 127.0.0.1] [--port 8765]
"""
from __future__ import annotations

import argparse
import io
import json
import os
import re
import shutil
import sys
from collections import Counter
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
from typing import Optional

from fastapi import FastAPI, HTTPException, Response
from fastapi.responses import JSONResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image
from pydantic import BaseModel

APP_DIR = Path(__file__).resolve().parent
CACHE_DIR = APP_DIR / ".cache"
SAMPLE_DIR = "표본결과"
DRAFT_DIR = "_draft"
STATUS_KEYS = ("일치", "보정성공", "미검출", "보정실패", "악화", "제외")
SUFFIX = {"answer": ".answer.json", "ao": ".aiocr.json", "harness": ".harness.json", "grade": ".grade.json"}

# create_app()이 한 번 채우는 전역 상태 (e2e 루트별 재바인딩 가능).
E2E: Path
PRE_DIR: Path
GS = BA = None


# ── e2e 경로 해석 / 데이터 경로 ───────────────────────────────────────────────
def default_e2e() -> Path:
    env = os.environ.get("LV_E2E")
    if env:
        return Path(env).resolve()
    for p in [APP_DIR, *APP_DIR.parents]:
        if (p / "e2e" / "grade_samples.py").is_file():
            return (p / "e2e").resolve()
    raise RuntimeError("e2e/grade_samples.py 를 찾을 수 없습니다 (--e2e 로 지정하세요)")

def sample_root() -> Path:
    return E2E / SAMPLE_DIR

def all_folders(folder: Optional[str] = None) -> list:
    if folder:
        return [folder]
    root = sample_root()
    return sorted(p.name for p in root.iterdir() if p.is_dir() and not p.name.startswith("_"))

def doc_files(folders: list):
    """folders 안의 (folder, file) 쌍을 grade.json 기준으로 순회한다."""
    for folder in folders:
        for gp in sorted((sample_root() / folder).glob("*.grade.json")):
            yield folder, gp.name[: -len(".grade.json")]

def spath(folder: str, file: str, kind: str) -> Optional[Path]:
    """draft/answer/ao/ao_ui/harness/grade 공용 경로 조립기."""
    if kind == "draft":
        return sample_root() / DRAFT_DIR / folder / f"{file}.draft.json"
    if kind == "ao_ui":
        cands = sorted((E2E / "out").glob("ao-ui-*"))
        return cands[-1] / folder / f"{file}.aiocr.ui.json" if cands else None
    suf = SUFFIX.get(kind)
    return sample_root() / folder / f"{file}{suf}" if suf else None

def load_json(p: Optional[Path]):
    if p is None or not p.exists():
        return None
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return None

def default_review() -> dict:
    return {"status": "", "checked": [], "edited": {}, "at": ""}

def check_doc(folder: str, file: str) -> Path:
    """(folder, file) 검증 + traversal 금지. 통과하면 원본 이미지 경로를 돌려준다."""
    if not folder or folder.startswith("_") or "/" in folder or "\\" in folder:
        raise HTTPException(404, "invalid folder")
    if not file or "/" in file or "\\" in file or file in (".", ".."):
        raise HTTPException(404, "invalid file")
    root = sample_root().resolve()
    fdir = (root / folder).resolve()
    if root not in fdir.parents or not fdir.is_dir():
        raise HTTPException(404, "folder not found")
    fpath = (fdir / file).resolve()
    if fdir not in fpath.parents or not fpath.is_file():
        raise HTTPException(404, "file not found")
    return fpath


# ── 상태 판정 (API.md 상태 표) ─────────────────────────────────────────────
def classify(key: str, ao_v: str, h_v: str, ao_val, h_val) -> str:
    if h_v == GS.SKIP:
        return "제외"
    ao_ok, h_ok = ao_v == GS.MATCH, h_v == GS.MATCH
    if ao_ok and h_ok:
        return "일치"
    if h_ok:
        return "보정성공"
    if ao_ok:
        return "악화"
    changed = (GS.canon(key, ao_val) or "") != (GS.canon(key, h_val) or "")
    return "보정실패" if changed else "미검출"

def not_dc(cells: list) -> list:
    return [c for c in cells if not (c.get("row") if isinstance(c, dict) else c.row).startswith("DC")]

def status_counts(cells: list) -> Counter:
    return Counter(classify(c.get("key"), c.get("ao_verdict"), c.get("h_verdict"), c.get("ao"), c.get("harness"))
                   for c in cells)


# ── AO UI(bbox) ──────────────────────────────────────────────────────────
def ui_result(folder: str, file: str):
    j = load_json(spath(folder, file, "ao_ui"))
    docs = (j or {}).get("documents") or []
    return docs[0].get("result") if docs else None

def ui_index(result) -> dict:
    idx: dict = {}
    if not result:
        return idx
    for c in result.get("fields") or []:
        idx[c["key"]] = c
    for g in result.get("groups") or []:
        for c in g.get("fields") or []:
            idx[c["key"]] = c
    for t in result.get("tables") or []:
        for row in t.get("rows") or []:
            for c in row:
                idx[c["key"]] = c
    return idx

def ui_key(area: str, container: str, row: str, key: str) -> Optional[str]:
    if area == "필드":
        return key
    if area == "그룹":
        return f"{container}.{key}"
    digits = re.sub(r"\D", "", row)
    return f"{container}[{int(digits) - 1}].{key}" if digits else None

def bbox_of(cell: dict) -> list:
    out = []
    for tb in cell.get("token_bbox") or []:
        for box in tb.get("token_bbox") or []:
            out.append({"page": tb.get("page"), "box": box})
    return out


# ── 하네스 evidence / path / editable ──────────────────────────────────────
def flatten_evidence(hb: dict) -> dict:
    ev = hb.get("evidence") or {}
    return {
        "tier": hb.get("tier") or "", "correction_basis": hb.get("correction_basis") or "",
        "decision_rule_no": hb.get("decision_rule_no") or "",
        "ao_value": hb.get("ao_value") or "", "final_value": hb.get("final_value") or "",
        "rule": ev.get("rule") or [], "reread": ev.get("reread") or {}, "master": ev.get("master") or {},
    }

def cell_path(area: str, container: str, row: str, key: str) -> str:
    if area == "필드":
        return f"fields[{key}]"
    if area == "그룹":
        return f"groups[{container}].fields[{key}]"
    digits = re.sub(r"\D", "", row)
    idx = int(digits) - 1 if digits else 0
    return f"tables[{container}].rows[{idx}].cells[{key}]"

def is_editable(area: str, row: str) -> bool:
    if row.startswith("AO") or row.startswith("결과"):
        return False
    return row.isdigit() if area == "표" else True


# ── 이미지 ────────────────────────────────────────────────────────────────
def page_count(path: Path) -> int:
    try:
        with Image.open(path) as im:
            return getattr(im, "n_frames", 1)
    except Exception:
        return 1

def render_png(path: Path, page: int) -> bytes:
    with Image.open(path) as im:
        if getattr(im, "n_frames", 1) > 1:
            im.seek(page - 1)
        im = im.convert("RGB")
        w, h = im.size
        m = max(w, h)
        if m > 3000:
            s = 3000 / m
            im = im.resize((max(1, int(w * s)), max(1, int(h * s))))
        buf = io.BytesIO()
        im.save(buf, "PNG")
        return buf.getvalue()

def cached_png(folder: str, file: str, page: int, src: Path) -> bytes:
    d = CACHE_DIR / folder
    d.mkdir(parents=True, exist_ok=True)
    cp = d / f"{file}.p{page}.png"
    if not cp.exists():
        cp.write_bytes(render_png(src, page))
    return cp.read_bytes()

def processed_path(folder: str, file: str, page: int) -> Optional[Path]:
    d = PRE_DIR / folder
    for name in (f"{file}.p{page}.png", f"{file}.png"):
        cp = d / name
        if cp.exists():
            return cp
    return None

def has_processed(folder: str, file: str, pages: int) -> bool:
    return any(processed_path(folder, file, p) for p in range(1, pages + 1))


# ── 문서 응답 조립 (GET/PUT /api/doc 공용) ─────────────────────────────────
def build_doc(folder: str, file: str, with_bbox: bool = True) -> dict:
    ans_path = spath(folder, file, "answer")
    if not ans_path.exists():
        raise HTTPException(404, "answer not found")
    cells, info = GS.grade_file(folder, ans_path)
    cells = not_dc(cells)
    doc_type = info.get("doc_type") or folder
    ans = GS.first_doc(GS.load(ans_path))
    hjson = GS.load(spath(folder, file, "harness"))
    R = GS.ResultDoc(GS.first_doc(hjson)) if hjson else None
    hmap: dict = {}
    if R is not None:
        def addh(area, container, row, label, key, answer, rc):
            if rc and rc.get("harness"):
                hmap[(area, container, str(row), key)] = flatten_evidence(rc["harness"])
        GS.walk(ans, R, doc_type, addh)
    uidx = ui_index(ui_result(folder, file)) if with_bbox else {}

    out_cells = []
    for c in cells:
        area, container, row, key = c.area, c.container, c.row, c.key
        cid = f"{area}|{container}|{row}|{key}"
        ukey = ui_key(area, container, row, key) if with_bbox else None
        uc = uidx.get(ukey) if ukey else None
        out_cells.append({
            "id": cid, "area": area, "container": container, "row": row, "label": c.label, "key": key,
            "answer": c.answer, "ao": c.ao, "harness": c.harness,
            "ao_verdict": c.ao_verdict, "h_verdict": c.h_verdict,
            "status": classify(key, c.ao_verdict, c.h_verdict, c.ao, c.harness),
            "kind": "누락" if c.h_verdict == GS.MISSING else "오탐" if c.h_verdict == GS.EXTRA else "",
            "effect": c.effect, "source": c.source, "tier": c.tier,
            "masked": c.masked, "uncertain": c.uncertain, "editable": is_editable(area, row),
            "confidence": uc.get("confidence") if uc else None,
            "bbox": bbox_of(uc) if uc else [],
            "evidence": hmap.get((area, container, row, key)),
            "path": cell_path(area, container, row, key),
        })

    draft = GS.load(spath(folder, file, "draft")) or {}
    review = draft.get("review") or default_review()
    meta = (GS.load(ans_path) or {}).get("meta") or {}
    t = GS.tally(cells)
    img_path = sample_root() / folder / file
    pages = page_count(img_path) if img_path.exists() else 1
    return {
        "folder": folder, "file": file, "doc_type": doc_type,
        "pages": pages, "processed": has_processed(folder, file, pages),
        "notes": meta.get("notes") or "", "uncertain": meta.get("uncertain") or [],
        "summary": {"채점칸": t["채점칸"], "AO정확도": t["AO정확도"], "하네스정확도": t["하네스정확도"],
                    "개선": t["개선"], "악화": t["악화"]},
        "harness_doc": (hjson and GS.first_doc(hjson).get("harness")) or {},
        "review": review,
        "cells": out_cells,
    }

def write_grade(folder: str, file: str, ans_path: Path):
    cells, info = GS.grade_file(folder, ans_path)
    t = GS.tally(cells)
    summary = {**{k: v for k, v in t.items() if k not in ("AO", "하네스")}, "AO": t["AO"], "하네스": t["하네스"]}
    spath(folder, file, "grade").write_text(
        json.dumps({"info": info, "summary": summary, "cells": [asdict(c) for c in cells]},
                   ensure_ascii=False, indent=1), encoding="utf-8")
    return cells, info


# ── draft 편집 ────────────────────────────────────────────────────────────
def _get_draft_value(draft: dict, area: str, container: str, row: str, key: str):
    if area == "필드":
        return (draft.get("fields") or {}).get(key)
    if area == "그룹":
        return (draft.get("groups", {}).get(container) or {}).get(key)
    headers = draft["tables"][container]["headers"]
    return draft["tables"][container]["rows"][int(row) - 1][headers.index(key)]

def _set_draft_value(draft: dict, area: str, container: str, row: str, key: str, value):
    if area == "필드":
        draft.setdefault("fields", {})[key] = value
    elif area == "그룹":
        draft.setdefault("groups", {}).setdefault(container, {})[key] = value
    else:
        headers = draft["tables"][container]["headers"]
        draft["tables"][container]["rows"][int(row) - 1][headers.index(key)] = value

def _add_result_row(draft: dict, cells: list, container: str, row: str, by: str):
    tbl = draft.setdefault("tables", {}).setdefault(container, {"headers": [], "rows": []})
    headers = tbl["headers"]
    field = "ao" if by == "ao" else "harness"
    matched = {c["key"]: c for c in cells
               if c["area"] == "표" and c["container"] == container and c["row"] == row}
    tbl["rows"].append([matched.get(h, {}).get(field) or "" for h in headers])


# ── 집계 (/api/stats, 엑셀 요약) ────────────────────────────────────────────
def _agg(cells: list, docs: int, done: int, counts: Optional[Counter] = None) -> dict:
    scored = (GS.MATCH, GS.MISMATCH, GS.MISSING, GS.EXTRA)
    n_ao = sum(1 for c in cells if c.get("ao_verdict") in scored)
    n_h = sum(1 for c in cells if c.get("h_verdict") in scored)
    ao_ok = sum(1 for c in cells if c.get("ao_verdict") == GS.MATCH)
    h_ok = sum(1 for c in cells if c.get("h_verdict") == GS.MATCH)
    counts = counts if counts is not None else status_counts(cells)
    return {
        "문서수": docs, "채점칸": n_ao,
        "AO정확도": ao_ok / n_ao if n_ao else None, "하네스정확도": h_ok / n_h if n_h else None,
        "보정성공": counts["보정성공"], "보정실패": counts["보정실패"], "악화": counts["악화"],
        "미검출": counts["미검출"],
        "누락": sum(1 for c in cells if c.get("h_verdict") == GS.MISSING),
        "오탐": sum(1 for c in cells if c.get("h_verdict") == GS.EXTRA),
        "하네스수정칸": sum(1 for c in cells if c.get("effect")),
        "검수완료문서": done,
    }

def compute_stats(folders: list) -> dict:
    out = {}
    all_cells, all_docs, all_done = [], 0, 0
    for f in folders:
        cells_f, docs, done = [], 0, 0
        for _, file in doc_files([f]):
            docs += 1
            cells_f.extend(not_dc((load_json(spath(f, file, "grade")) or {}).get("cells") or []))
            review = (load_json(spath(f, file, "draft")) or {}).get("review") or {}
            if review.get("status") == "done":
                done += 1
        out[f] = _agg(cells_f, docs, done)
        all_cells.extend(cells_f)
        all_docs += docs
        all_done += done
    out["전체"] = _agg(all_cells, all_docs, all_done)
    return out

def evidence_summary(cell: dict) -> str:
    ev = cell.get("evidence")
    if not ev:
        return ""
    bits = [f"tier={ev.get('tier') or ''}" + (f" basis={ev['correction_basis']}" if ev.get("correction_basis") else "")]
    for r in ev.get("rule") or []:
        if r.get("result") in ("fail", "warn"):
            bits.append(f"{r.get('rule_id')}:{r.get('result')} {r.get('detail') or ''}".strip())
    rr = ev.get("reread") or {}
    if rr.get("status") not in (None, "", "unavailable"):
        tag = f"({rr['engine_id']})" if rr.get("engine_id") and rr["engine_id"] != "none" else ""
        bits.append(f"reread:{rr['status']}={rr.get('value')}{tag}")
    m = ev.get("master") or {}
    if m.get("status"):
        bits.append(f"master:{m['status']}")
    return " | ".join(bits)[:32000]


# ── 엑셀 export ───────────────────────────────────────────────────────────
def build_workbook(folders: list) -> Workbook:
    stats = compute_stats(folders)
    wb = Workbook()
    wb.remove(wb.active)
    HEAD, HEAD_FONT = PatternFill("solid", fgColor="1F4E78"), Font(bold=True, color="FFFFFF")
    STATUS_FILL = {"일치": "E2EFDA", "보정성공": "C6EFCE", "미검출": "FCE4D6",
                   "보정실패": "F8CBAD", "악화": "FFC7CE", "제외": "EDEDED"}

    def sheet(title, headers, rows, widths=None, color_col=None):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.fill, c.font = HEAD, HEAD_FONT
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        for i in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 14)
        if color_col:
            fills = {k: PatternFill("solid", fgColor=v) for k, v in STATUS_FILL.items()}
            for (cell,) in ws.iter_rows(min_row=2, min_col=color_col, max_col=color_col):
                if cell.value in fills:
                    cell.fill = fills[cell.value]

    stat_cols = ["문서수", "채점칸", "AO정확도", "하네스정확도", "보정성공", "보정실패", "악화",
                 "미검출", "누락", "오탐", "하네스수정칸", "검수완료문서"]
    sheet("요약", ["구분", *stat_cols], [[k, *[v.get(c) for c in stat_cols]] for k, v in stats.items()])

    doc_rows, review_rows = [], []
    for folder, file in doc_files(folders):
        doc = build_doc(folder, file, with_bbox=False)
        review = doc["review"]
        counts = status_counts(doc["cells"])
        agg = _agg(doc["cells"], 1, 1 if review.get("status") == "done" else 0, counts)
        doc_rows.append([folder, file, doc["doc_type"], agg["채점칸"], agg["AO정확도"], agg["하네스정확도"],
                          *[counts[k] for k in STATUS_KEYS], agg["누락"], agg["오탐"],
                          review.get("status", ""), len(review.get("checked") or []),
                          len(review.get("edited") or {})])
        edited = review.get("edited") or {}
        checked = set(review.get("checked") or [])
        for c in doc["cells"]:
            review_rows.append([doc["doc_type"], file, c["area"], c["container"], c["row"], c["label"], c["key"],
                                 c["answer"], c["ao"], c["harness"], c["ao_verdict"], c["h_verdict"],
                                 c["status"], c["kind"], "Y" if c["id"] in edited else "",
                                 (edited.get(c["id"]) or {}).get("by", ""),
                                 "Y" if c["id"] in checked else "", evidence_summary(c)])
    sheet("문서별", ["폴더", "파일", "서식", "채점칸", "AO정확도", "하네스정확도",
                   *STATUS_KEYS, "누락", "오탐", "검수상태", "검수확인칸수", "수정칸수"],
          doc_rows, {2: 36})
    sheet("검수결과", ["서식", "파일", "구역", "그룹/표", "행", "행라벨", "필드", "Golden", "AO", "Harness",
                    "AO↔Golden", "Harness↔Golden", "상태", "유형", "수정여부", "수정출처", "검수확인", "Evidence"],
          review_rows, {2: 36, 6: 24, 18: 60}, color_col=13)

    ws = wb.create_sheet("RawJSON")
    ws.append(["파일", "종류", "JSON"])
    for c in ws[1]:
        c.fill, c.font = HEAD, HEAD_FONT
    for folder, file in doc_files(folders):
        for kind in ("answer", "ao", "harness"):
            j = load_json(spath(folder, file, kind))
            if j is None:
                continue
            s = json.dumps(j, ensure_ascii=False)
            for i in range(0, len(s), 32000):
                ws.append([file if i == 0 else "", kind if i == 0 else "", s[i:i + 32000]])
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["C"].width = 100
    return wb


# ── 요청 모델 ─────────────────────────────────────────────────────────────
class EditItem(BaseModel):
    id: str
    value: Optional[str] = None
    by: str = "manual"

class AddRow(BaseModel):
    container: str
    row: str
    by: str = "ao"

class ReviewPatch(BaseModel):
    status: Optional[str] = None
    checked: Optional[list] = None

class DocPatch(BaseModel):
    edits: list[EditItem] = []
    add_rows: list[AddRow] = []
    review: Optional[ReviewPatch] = None


# ── FastAPI 앱 ────────────────────────────────────────────────────────────
def create_app(e2e: Path, pre_dir: Path, static_dir: Optional[Path] = None) -> FastAPI:
    global E2E, PRE_DIR, GS, BA
    E2E, PRE_DIR = e2e.resolve(), pre_dir
    if str(E2E) not in sys.path:
        sys.path.insert(0, str(E2E))
    import build_answers as BA  # noqa: E402
    import grade_samples as GS  # noqa: E402

    app = FastAPI(title="Golden Set Viewer")

    @app.get("/api/docs")
    def api_docs():
        out = []
        for folder, file in doc_files(all_folders()):
            gj = load_json(spath(folder, file, "grade")) or {}
            info, summary = gj.get("info") or {}, gj.get("summary") or {}
            cells = not_dc(gj.get("cells") or [])
            counts = status_counts(cells)
            review = (load_json(spath(folder, file, "draft")) or {}).get("review") or {}
            out.append({
                "folder": folder, "file": file, "doc_type": info.get("doc_type") or folder,
                "ao_acc": summary.get("AO정확도"), "h_acc": summary.get("하네스정확도"),
                "counts": {k: counts.get(k, 0) for k in STATUS_KEYS},
                "errors": counts.get("미검출", 0) + counts.get("보정실패", 0) + counts.get("악화", 0),
                "review": review.get("status", ""),
                "checked": len(review.get("checked") or []),
                "edited": len(review.get("edited") or {}),
            })
        return out

    @app.get("/api/doc/{folder}/{file}")
    def api_doc(folder: str, file: str):
        check_doc(folder, file)
        return build_doc(folder, file)

    @app.get("/api/image/{folder}/{file}")
    def api_image(folder: str, file: str, view: str = "original", page: int = 1):
        src = check_doc(folder, file)
        if view == "processed":
            pp = processed_path(folder, file, page)
            if pp is not None:
                return Response(pp.read_bytes(), media_type="image/png", headers={"X-Image-View": "processed"})
            data = cached_png(folder, file, page, src)
            return Response(data, media_type="image/png",
                             headers={"X-Image-View": "original", "X-Image-Fell-Back": "true"})
        data = cached_png(folder, file, page, src)
        return Response(data, media_type="image/png", headers={"X-Image-View": "original"})

    @app.get("/api/raw/{folder}/{file}/{kind}")
    def api_raw(folder: str, file: str, kind: str):
        check_doc(folder, file)
        p = spath(folder, file, kind)
        if p is None or not p.exists():
            raise HTTPException(404, f"{kind} not found")
        return JSONResponse(load_json(p))

    @app.get("/api/stats")
    def api_stats(folder: Optional[str] = None):
        return compute_stats(all_folders(folder))

    @app.get("/api/export.xlsx")
    def api_export(folder: Optional[str] = None):
        wb = build_workbook(all_folders(folder))
        buf = io.BytesIO()
        wb.save(buf)
        return Response(buf.getvalue(),
                         media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                         headers={"Content-Disposition": "attachment; filename=golden_export.xlsx"})

    @app.put("/api/doc/{folder}/{file}")
    def api_put(folder: str, file: str, patch: DocPatch):
        check_doc(folder, file)
        ans_path = spath(folder, file, "answer")
        dpath = spath(folder, file, "draft")
        if not ans_path.exists() or not dpath.exists():
            raise HTTPException(404, "draft/answer not found")

        pre_cells, _ = GS.grade_file(folder, ans_path)
        pre_rows = [asdict(c) for c in pre_cells]

        backup = sample_root() / DRAFT_DIR / "_backup" / folder / f"{file}.draft.json"
        if not backup.exists():
            backup.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(dpath, backup)

        draft = json.loads(dpath.read_text(encoding="utf-8"))
        review = draft.get("review") or default_review()
        review.setdefault("edited", {})
        review.setdefault("checked", [])
        now = datetime.now().isoformat(timespec="seconds")

        for e in patch.edits:
            area, container, row, key = e.id.split("|", 3)
            if not is_editable(area, row):
                raise HTTPException(400, f"editable 아님: {e.id}")
            old = _get_draft_value(draft, area, container, row, key)
            _set_draft_value(draft, area, container, row, key, e.value)
            review["edited"][e.id] = {"from": old, "to": e.value, "by": e.by, "at": now}

        for ar in patch.add_rows:
            _add_result_row(draft, pre_rows, ar.container, ar.row, ar.by)

        if patch.review is not None:
            if patch.review.status is not None:
                review["status"] = patch.review.status
            if patch.review.checked is not None:
                review["checked"] = patch.review.checked
        review["at"] = now
        draft["review"] = review
        dpath.write_text(json.dumps(draft, ensure_ascii=False, indent=2), encoding="utf-8")

        answer = BA.convert(draft)
        answer["meta"]["review"] = review
        ans_path.write_text(json.dumps(answer, ensure_ascii=False, indent=2), encoding="utf-8")

        write_grade(folder, file, ans_path)
        return build_doc(folder, file)

    if static_dir and static_dir.exists():
        app.mount("/", StaticFiles(directory=str(static_dir), html=True), name="static")
    return app


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--e2e", type=Path, default=None, help="e2e 루트 경로")
    ap.add_argument("--pre-dir", type=Path, default=None, help="전처리 이미지 경로")
    ap.add_argument("--host", default="127.0.0.1")
    ap.add_argument("--port", type=int, default=8765)
    args = ap.parse_args()

    e2e = args.e2e.resolve() if args.e2e else default_e2e()
    pre_dir = args.pre_dir.resolve() if args.pre_dir else (e2e / "out" / "ao-pre-image")
    app = create_app(e2e, pre_dir, APP_DIR / "static")

    import uvicorn
    uvicorn.run(app, host=args.host, port=args.port)


if __name__ == "__main__":
    main()
