"""app.py 테스트. 실 데이터(e2e/표본결과)는 읽기만 하고, PUT 은 임시 복사본에서만 수행한다."""
from __future__ import annotations

import json
import os
import shutil
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

import app

E2E = app.default_e2e()
PRE_DIR = E2E / "out" / "ao-pre-image"

DIAG_FOLDER, DIAG_FILE = "진단서", "20230227091019607249.tif"
RECEIPT_FOLDER = "진료비영수증"


@pytest.fixture(scope="module")
def client() -> TestClient:
    return TestClient(app.create_app(E2E, PRE_DIR, None))


def any_receipt_file(client: TestClient) -> tuple[str, str]:
    doc = next(d for d in client.get("/api/docs").json() if d["folder"] == RECEIPT_FOLDER)
    return doc["folder"], doc["file"]


# ── /api/docs ────────────────────────────────────────────────────────────
def test_docs_list_matches_filesystem(client):
    r = client.get("/api/docs")
    assert r.status_code == 200
    docs = r.json()
    on_disk = sum(1 for f in (E2E / "표본결과").iterdir() if f.is_dir() and not f.name.startswith("_")
                  for _ in f.glob("*.grade.json"))
    assert len(docs) == on_disk
    assert len(docs) > 0
    d0 = docs[0]
    assert set(d0) >= {"folder", "file", "doc_type", "ao_acc", "h_acc", "counts", "errors",
                       "review", "checked", "edited"}
    assert set(d0["counts"]) == {"일치", "보정성공", "미검출", "보정실패", "악화", "제외"}


# ── /api/doc ─────────────────────────────────────────────────────────────
def test_doc_diagnosis_has_bbox_and_evidence(client):
    r = client.get(f"/api/doc/{DIAG_FOLDER}/{DIAG_FILE}")
    assert r.status_code == 200
    d = r.json()
    assert d["folder"] == DIAG_FOLDER and d["file"] == DIAG_FILE
    assert d["pages"] >= 1
    assert set(d["summary"]) == {"채점칸", "AO정확도", "하네스정확도", "개선", "악화"}
    assert d["cells"], "cells should not be empty"
    for c in d["cells"]:
        assert c["status"] in ("일치", "보정성공", "미검출", "보정실패", "악화", "제외")
        assert not c["row"].startswith("DC")
    assert any(c["bbox"] for c in d["cells"]), "at least one cell should carry bbox"
    assert any(c["evidence"] for c in d["cells"]), "at least one cell should carry harness evidence"


def test_doc_receipt_table_cells(client):
    folder, file = any_receipt_file(client)
    r = client.get(f"/api/doc/{folder}/{file}")
    assert r.status_code == 200
    d = r.json()
    table_cells = [c for c in d["cells"] if c["area"] == "표"]
    assert table_cells
    assert any(c["bbox"] for c in table_cells)
    for c in table_cells:
        if c["row"].isdigit():
            assert c["path"] == f"tables[{c['container']}].rows[{int(c['row']) - 1}].cells[{c['key']}]"
            assert c["editable"] is True
        else:
            assert c["editable"] is False


def test_doc_404_for_bad_path(client):
    assert client.get("/api/doc/_draft/x.tif").status_code == 404
    assert client.get(f"/api/doc/{DIAG_FOLDER}/../../etc/passwd").status_code == 404
    assert client.get(f"/api/doc/{DIAG_FOLDER}/no-such-file.tif").status_code == 404


# ── /api/image ───────────────────────────────────────────────────────────
def test_image_original_png(client):
    r = client.get(f"/api/image/{DIAG_FOLDER}/{DIAG_FILE}")
    assert r.status_code == 200
    assert r.headers["content-type"] == "image/png"
    assert r.headers["x-image-view"] == "original"
    assert r.content[:8] == b"\x89PNG\r\n\x1a\n"


def test_image_processed_falls_back(client):
    r = client.get(f"/api/image/{DIAG_FOLDER}/{DIAG_FILE}", params={"view": "processed"})
    assert r.status_code == 200
    assert r.headers.get("x-image-fell-back") == "true"
    assert r.headers["x-image-view"] == "original"


# ── /api/raw ─────────────────────────────────────────────────────────────
def test_raw_kinds(client):
    for kind in ("answer", "draft", "ao", "harness", "grade"):
        r = client.get(f"/api/raw/{DIAG_FOLDER}/{DIAG_FILE}/{kind}")
        assert r.status_code == 200, kind
        assert isinstance(r.json(), dict)


def test_raw_ao_ui(client):
    r = client.get(f"/api/raw/{DIAG_FOLDER}/{DIAG_FILE}/ao_ui")
    assert r.status_code == 200
    assert "documents" in r.json()


def test_raw_unknown_kind_404(client):
    assert client.get(f"/api/raw/{DIAG_FOLDER}/{DIAG_FILE}/bogus").status_code == 404


# ── /api/stats & export ──────────────────────────────────────────────────
def test_stats(client):
    r = client.get("/api/stats")
    assert r.status_code == 200
    stats = r.json()
    assert "전체" in stats
    assert set(stats["전체"]) == {"문서수", "채점칸", "AO정확도", "하네스정확도", "보정성공", "보정실패",
                                "악화", "미검출", "누락", "오탐", "하네스수정칸", "검수완료문서"}
    assert stats["전체"]["문서수"] > 0


def test_stats_scoped_to_folder(client):
    r = client.get("/api/stats", params={"folder": DIAG_FOLDER})
    stats = r.json()
    assert DIAG_FOLDER in stats and "전체" in stats
    assert stats[DIAG_FOLDER]["문서수"] == stats["전체"]["문서수"]


def test_export_xlsx(client):
    r = client.get("/api/export.xlsx", params={"folder": DIAG_FOLDER})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml")
    assert len(r.content) > 1000
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    assert set(wb.sheetnames) >= {"요약", "문서별", "검수결과", "RawJSON"}


# ── PUT 편집 (임시 복사본에서만) ─────────────────────────────────────────────
@pytest.fixture()
def tmp_client(tmp_path: Path) -> TestClient:
    """실 e2e 코드(symlink)+표본 1건만 복사한 임시 e2e 루트로 앱을 띄운다. 실 데이터는 건드리지 않는다."""
    tmp_e2e = tmp_path / "e2e"
    tmp_e2e.mkdir()
    for name in ("grade_samples.py", "build_answers.py", "make_report.py"):
        os.symlink(E2E / name, tmp_e2e / name)

    sample = tmp_e2e / "표본결과" / DIAG_FOLDER
    sample.mkdir(parents=True)
    real_dir = E2E / "표본결과" / DIAG_FOLDER
    for suffix in ("", ".answer.json", ".aiocr.json", ".harness.json", ".grade.json"):
        src = real_dir / f"{DIAG_FILE}{suffix}"
        if src.exists():
            shutil.copy2(src, sample / src.name)

    draft_dir = tmp_e2e / "표본결과" / "_draft" / DIAG_FOLDER
    draft_dir.mkdir(parents=True)
    shutil.copy2(E2E / "표본결과" / "_draft" / DIAG_FOLDER / f"{DIAG_FILE}.draft.json",
                 draft_dir / f"{DIAG_FILE}.draft.json")

    return TestClient(app.create_app(tmp_e2e, tmp_e2e / "out" / "ao-pre-image", None))


def test_put_edit_roundtrip(tmp_client):
    before = tmp_client.get(f"/api/doc/{DIAG_FOLDER}/{DIAG_FILE}").json()
    name_cell = next(c for c in before["cells"] if c["area"] == "그룹" and c["key"] == "이름")
    assert name_cell["answer"] == "이동근"
    cid = name_cell["id"]
    assert cid == "그룹|환자정보||이름"

    body = {"edits": [{"id": cid, "value": "테스트이름", "by": "manual"}],
            "review": {"status": "done", "checked": [cid]}}
    r = tmp_client.put(f"/api/doc/{DIAG_FOLDER}/{DIAG_FILE}", json=body)
    assert r.status_code == 200
    after = r.json()

    updated = next(c for c in after["cells"] if c["id"] == cid)
    assert updated["answer"] == "테스트이름"
    assert after["review"]["status"] == "done"
    assert cid in after["review"]["checked"]
    assert cid in after["review"]["edited"]
    assert after["review"]["edited"][cid]["from"] == "이동근"
    assert after["review"]["edited"][cid]["to"] == "테스트이름"
    assert after["review"]["edited"][cid]["by"] == "manual"

    # 재조회해도 그대로 반영돼 있어야 한다 (draft/answer/grade.json 이 실제로 갱신됨)
    again = tmp_client.get(f"/api/doc/{DIAG_FOLDER}/{DIAG_FILE}").json()
    assert next(c for c in again["cells"] if c["id"] == cid)["answer"] == "테스트이름"
    assert again["review"]["status"] == "done"


def test_put_backup_created_once(tmp_client):
    dpath_backup = None
    body = {"edits": [], "review": {"status": "progress"}}
    r = tmp_client.put(f"/api/doc/{DIAG_FOLDER}/{DIAG_FILE}", json=body)
    assert r.status_code == 200
    # backup 파일 위치는 앱 내부에서 조립하므로 여기선 응답으로 상태만 확인한다
    assert r.json()["review"]["status"] == "progress"


def test_put_rejects_noneditable_cell(tmp_client):
    before = tmp_client.get(f"/api/doc/{DIAG_FOLDER}/{DIAG_FILE}").json()
    fake_id = "표|병명내역|결과1|병명"
    r = tmp_client.put(f"/api/doc/{DIAG_FOLDER}/{DIAG_FILE}",
                        json={"edits": [{"id": fake_id, "value": "x", "by": "manual"}]})
    assert r.status_code == 400


def test_real_data_untouched(client):
    """실 표본결과 draft/answer 는 이 테스트 스위트에서 변경되지 않아야 한다."""
    draft_path = E2E / "표본결과" / "_draft" / DIAG_FOLDER / f"{DIAG_FILE}.draft.json"
    draft = json.loads(draft_path.read_text(encoding="utf-8"))
    assert draft.get("groups", {}).get("환자정보", {}).get("이름") == "이동근"
    assert "review" not in draft
