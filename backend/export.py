"""번들 내보내기: 전체 묶음 ZIP, Golden Excel."""
from __future__ import annotations

import io
import tempfile
import zipfile
from pathlib import Path
from typing import BinaryIO, Callable

from openpyxl import Workbook
from openpyxl.cell.cell import Cell, ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill

from . import bundle as B

_STATUS_FILL = {
    "MATCH": PatternFill("solid", fgColor="C6EFCE"),
    "MISMATCH": PatternFill("solid", fgColor="FFC7CE"),
    "MISSING": PatternFill("solid", fgColor="FFD9A0"),
    "EXTRA": PatternFill("solid", fgColor="E4C7FF"),
    "TYPE_MISMATCH": PatternFill("solid", fgColor="FFF2A8"),
}


_STORED_EXTS = {".png", ".jpg", ".jpeg", ".webp"}  # 이미 압축된 이미지는 다시 deflate 하지 않는다


class ExportCancelled(Exception):
    """Raised when a background export has received a cancellation request."""


def _checkpoint(cancel_check: Callable[[], bool] | None) -> None:
    if cancel_check and cancel_check():
        raise ExportCancelled("export cancelled")


def _progress(progress: Callable[[int, int, str], None] | None,
              completed: int, total: int, phase: str) -> None:
    if progress:
        progress(completed, total, phase)


def _resolve_doc_ids(data_dir: Path, bundle_id: str, doc_id: str | None, scope: str,
                     ids: list[str] | None, *, state_snapshot: dict | None = None) -> tuple[Path, list[str]]:
    bdir = B.bundle_dir(data_dir, bundle_id)
    if not bdir.is_dir():
        raise B.ApiError(404, "bundle not found")
    if doc_id is not None and ids is not None:
        raise B.ApiError(422, "doc_id and ids cannot be combined")
    all_ids = B.doc_ids(bdir)
    if doc_id is not None:
        B._safe_id(doc_id, "doc_id")
        if doc_id not in all_ids:
            raise B.ApiError(404, "document not found")
        return bdir, [doc_id]
    if ids is not None:
        if not isinstance(ids, list) or not ids:
            raise B.ApiError(422, "ids must be a non-empty list")
        if not all(isinstance(item, str) for item in ids):
            raise B.ApiError(422, "ids must contain strings")
        for item in ids:
            B._safe_id(item, "doc_id")
        unknown = set(ids) - set(all_ids)
        if unknown:
            raise B.ApiError(404, f"document not found: {sorted(unknown)[:5]}")
        selected = set(ids)
        return bdir, [item for item in all_ids if item in selected]
    return bdir, B.scope_doc_ids(data_dir, bundle_id, scope, state_snapshot=state_snapshot)


def export_bundle_zip(data_dir: Path, bundle_id: str, doc_id: str | None = None,
                      scope: str = "enabled", *, ids: list[str] | None = None,
                      progress: Callable[[int, int, str], None] | None = None,
                      cancel_check: Callable[[], bool] | None = None,
                      out: BinaryIO | None = None) -> BinaryIO | None:
    """원본·전처리 이미지와 AO·Harness·Golden JSON을 업로드 폴더 구조 그대로 묶는다.
    doc_id가 있으면 그 문서만, 없으면 scope(활성·비활성) 문서만. 수천 건이면 수 GB라 메모리 대신 임시 파일에 쓴다."""
    bdir, selected_ids = _resolve_doc_ids(data_dir, bundle_id, doc_id, scope, ids)
    wanted = set(selected_ids)
    files_by_doc: dict[str, list[tuple[str, Path]]] = {did: [] for did in selected_ids}
    for kind in (*B.DOC_KINDS, "ao_ui"):
        d = bdir / kind
        if not d.is_dir():
            continue
        for p in sorted(d.iterdir()):
            if p.is_file() and not p.name.startswith(".") and p.stem in wanted:
                files_by_doc[p.stem].append((kind, p))
    tmp_dir = B.cache_root(data_dir) / "tmp"
    tmp_dir.mkdir(parents=True, exist_ok=True)
    owned = out is None
    if out is None:
        out = tempfile.TemporaryFile(dir=tmp_dir)
    try:
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
            total = len(selected_ids)
            _progress(progress, 0, total, "preparing")
            for completed, did in enumerate(selected_ids, 1):
                _checkpoint(cancel_check)
                for kind, p in files_by_doc[did]:
                    _checkpoint(cancel_check)
                    ctype = zipfile.ZIP_STORED if p.suffix.lower() in _STORED_EXTS else zipfile.ZIP_DEFLATED
                    info = zipfile.ZipInfo(f"{kind}/{p.name}")
                    info.compress_type = ctype
                    with p.open("rb") as src, zf.open(info, "w", force_zip64=True) as dst:
                        while True:
                            _checkpoint(cancel_check)
                            block = src.read(1 << 20)
                            if not block:
                                break
                            dst.write(block)
                _progress(progress, completed, total, "writing")
            _progress(progress, total, total, "finalizing")
            _checkpoint(cancel_check)
        _checkpoint(cancel_check)
    except BaseException:
        if owned:
            out.close()
        raise
    out.seek(0)
    return out if owned else None


_OX = {True: "O", False: "X", None: ""}


def _cell_val(v) -> str:
    return "" if v is None else ILLEGAL_CHARACTERS_RE.sub("", str(v))  # OCR 값의 제어문자는 xlsx에 쓸 수 없다



def _append_row(ws, values) -> None:
    # OCR 텍스트와 문서 ID는 '='로 시작해도 수식으로 실행하지 않는다.
    cells = []
    for value in values:
        if isinstance(value, str):
            cell = Cell(ws, value=_cell_val(value))
            cell.data_type = "s"
            cells.append(cell)
        else:
            cells.append(value)
    ws.append(cells)


def export_golden_xlsx(data_dir: Path, bundle_id: str, doc_id: str | None = None, scope: str = "enabled", *,
                       ids: list[str] | None = None,
                       progress: Callable[[int, int, str], None] | None = None,
                       cancel_check: Callable[[], bool] | None = None,
                       out: BinaryIO | None = None) -> bytes | None:
    bdir = B.bundle_dir(data_dir, bundle_id)
    if not bdir.is_dir():
        raise B.ApiError(404, "bundle not found")
    state_snapshot = B.load_state(bdir)
    _, selected_ids = _resolve_doc_ids(data_dir, bundle_id, doc_id, scope, ids,
                                       state_snapshot=state_snapshot)

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "요약"
    _append_row(ws_summary, ["id", "doc_type", "AO 분류", "H 분류", "검수상태",
                        "AO MATCH", "AO MISMATCH", "AO MISSING", "AO EXTRA", "AO TYPE_MISMATCH", "AO 정확도",
                        "H MATCH", "H MISMATCH", "H MISSING", "H EXTRA", "H TYPE_MISMATCH", "H 정확도"])
    for c in ws_summary[1]:
        c.font = Font(bold=True)

    ws_fields = wb.create_sheet("필드")
    _append_row(ws_fields, ["문서", "문서index", "구역", "그룹", "key", "value", "dtype"])
    for c in ws_fields[1]:
        c.font = Font(bold=True)

    ws_tables = wb.create_sheet("표")

    ws_compare = wb.create_sheet("비교")
    _append_row(ws_compare, ["문서", "path", "구역", "그룹/표", "행", "key", "Golden", "AO", "AO상태", "Harness", "Harness상태"])
    for c in ws_compare[1]:
        c.font = Font(bold=True)

    totals = {"ao": {"MATCH": 0, "MISMATCH": 0, "MISSING": 0, "EXTRA": 0, "TYPE_MISMATCH": 0},
              "h": {"MATCH": 0, "MISMATCH": 0, "MISSING": 0, "EXTRA": 0, "TYPE_MISMATCH": 0}}

    compare_row = 1
    table_row = 0
    total = len(selected_ids)
    _progress(progress, 0, total, "preparing")
    for completed, did in enumerate(selected_ids, 1):
        _checkpoint(cancel_check)
        detail = B.doc_detail(data_dir, bundle_id, did, state_snapshot=state_snapshot,
                              include_pages=False)
        ao_sc, h_sc = detail["score"]["ao"], detail["score"]["harness"]
        cls = detail["classification"]
        row = [did, detail["doc_type"], *(_OX[cls[k]] for k in ("ao", "harness")), detail["review"] or "",
               *(ao_sc.get(k) if ao_sc else "" for k in ("MATCH", "MISMATCH", "MISSING", "EXTRA", "TYPE_MISMATCH")),
               ao_sc.get("accuracy") if ao_sc else "",
               *(h_sc.get(k) if h_sc else "" for k in ("MATCH", "MISMATCH", "MISSING", "EXTRA", "TYPE_MISMATCH")),
               h_sc.get("accuracy") if h_sc else ""]
        _append_row(ws_summary, row)
        for side, sc, ok in (("ao", ao_sc, cls["ao"]), ("h", h_sc, cls["harness"])):
            if sc and ok is not False:
                for k in ("MATCH", "MISMATCH", "MISSING", "EXTRA", "TYPE_MISMATCH"):
                    totals[side][k] += sc[k]

        golden = detail["golden"] or {}
        for doc_i, doc in enumerate(golden.get("documents") or []):
            for c in doc.get("extracted_fields") or []:
                _checkpoint(cancel_check)
                _append_row(ws_fields, [did, doc_i, "필드", "", c.get("key"), _cell_val(c.get("value")), c.get("dtype", "")])
            for g in doc.get("extracted_groups") or []:
                for c in g.get("fields") or []:
                    _checkpoint(cancel_check)
                    _append_row(ws_fields, [did, doc_i, "그룹", g.get("key"), c.get("key"), _cell_val(c.get("value")),
                                       c.get("dtype", "")])
            for t in doc.get("extracted_tables") or []:
                _checkpoint(cancel_check)
                headers = t.get("headers") or []
                _append_row(ws_tables, [did, t.get("key")] + headers)
                table_row += 1
                for column in range(1, len(headers) + 3):
                    ws_tables.cell(row=table_row, column=column).font = Font(bold=True)
                for r in t.get("rows") or []:
                    _checkpoint(cancel_check)
                    cellmap = {c.get("key"): c for c in r}
                    _append_row(ws_tables, ["", ""] + [_cell_val((cellmap.get(h) or {}).get("value")) for h in headers])
                table_row += len(t.get("rows") or []) + 1
                _append_row(ws_tables, [])

        for r in detail["compare"]:
            _checkpoint(cancel_check)
            _append_row(ws_compare, [did, r["path"], r["area"], r["container"], r["row"], r["key"],
                                _cell_val(r["golden"]), _cell_val(r["ao"]), r["ao_status"],
                                _cell_val(r["harness"]), r["harness_status"]])
            compare_row += 1
            for col_idx, status in ((9, r["ao_status"]), (11, r["harness_status"])):
                fill = _STATUS_FILL.get(status)
                if fill:
                    ws_compare.cell(row=compare_row, column=col_idx).fill = fill
        _progress(progress, completed, total, "writing")

    def _acc(n_match: int, total: dict) -> float | str:
        t = sum(total.values())
        return round(n_match / t, 4) if t else ""

    _append_row(ws_summary, ["합계", "", "", "", "",
                        totals["ao"]["MATCH"], totals["ao"]["MISMATCH"], totals["ao"]["MISSING"],
                        totals["ao"]["EXTRA"], totals["ao"]["TYPE_MISMATCH"], _acc(totals["ao"]["MATCH"], totals["ao"]),
                        totals["h"]["MATCH"], totals["h"]["MISMATCH"], totals["h"]["MISSING"],
                        totals["h"]["EXTRA"], totals["h"]["TYPE_MISMATCH"], _acc(totals["h"]["MATCH"], totals["h"])])
    for c in ws_summary[ws_summary.max_row]:
        c.font = Font(bold=True)

    _progress(progress, total, total, "finalizing")
    _checkpoint(cancel_check)
    if out is not None:
        wb.save(out)
        _checkpoint(cancel_check)
        return None
    buf = io.BytesIO()
    wb.save(buf)
    _checkpoint(cancel_check)
    return buf.getvalue()
