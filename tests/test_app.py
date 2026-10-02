"""backend 통합 테스트. 더미 번들 생성기로 tmp_path 에 데이터를 만들고 FastAPI TestClient 로 검증한다."""
from __future__ import annotations

import hashlib
import io
import json
import sys
import zipfile
from urllib.parse import quote
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from backend.app import create_app
from backend.bundle import classify, stem_of
from backend.doctype import DOC_TYPES, TEMPLATES, apply_template, canon, label
from scripts import make_dummy_bundle as dummy


def _entries(root: Path) -> list[tuple[str, bytes]]:
    return [(str(p.relative_to(root)), p.read_bytes()) for p in sorted(root.rglob("*")) if p.is_file()]


@pytest.fixture
def dummy_root(tmp_path) -> Path:
    root = tmp_path / "dummy_src"
    root.mkdir()
    dummy.build(root)
    return root


@pytest.fixture
def client(tmp_path) -> TestClient:
    app = create_app(tmp_path / "storage")
    return TestClient(app)


def _upload_folder(client: TestClient, root: Path, name: str = "dummy-folder"):
    files = [("files", (relpath, data, "application/octet-stream")) for relpath, data in _entries(root)]
    return client.post("/api/bundles", files=files, data={"name": name})


@pytest.fixture
def bundle(client: TestClient, dummy_root: Path) -> dict:
    resp = _upload_folder(client, dummy_root)
    assert resp.status_code == 201, resp.text
    return resp.json()


# ── health ────────────────────────────────────────────────────────────────

def test_health(client: TestClient):
    r = client.get("/api/health")
    assert r.status_code == 200
    assert r.json() == {"ok": True}


# ── classify / stem ──────────────────────────────────────────────────────

def test_classify_folders():
    assert classify("original/MC001.png") == "original"
    assert classify("preprocessed/x.png") == "preprocessed"
    assert classify("ao_extract/x.aiocr.json") == "ao_extract"
    assert classify("harness/x.harness.json") == "harness"
    assert classify("golden/x.json") == "golden"
    assert classify("정답지/x.json") == "golden"
    assert classify("foo/__MACOSX/bar.png") is None


def test_classify_suffix_fallback():
    assert classify("flat/ABC001.answer.json") == "golden"
    assert classify("flat/ABC001.harness.json") == "harness"
    assert classify("flat/ABC001.aiocr.json") == "ao_extract"
    assert classify("flat/ABC001.png") == "original"
    assert classify("flat/ABC001.unrelated.json") is None


def test_stem_across_extensions():
    assert stem_of("ABC001.jpg") == "ABC001"
    assert stem_of("ABC001.png") == "ABC001"
    assert stem_of("ABC001.json") == "ABC001"
    assert stem_of("ABC001.tif.aiocr.json") == "ABC001"
    assert stem_of("ABC001.p2.png") == "ABC001"
    assert stem_of("ABC001.answer.json") == "ABC001"


# ── upload ────────────────────────────────────────────────────────────────

def test_upload_zip(client: TestClient, dummy_root: Path):
    zip_bytes = io.BytesIO()
    with zipfile.ZipFile(zip_bytes, "w") as zf:
        for relpath, data in _entries(dummy_root):
            zf.writestr(relpath, data)
    resp = client.post("/api/bundles", files={"files": ("dummy.zip", zip_bytes.getvalue(), "application/zip")})
    assert resp.status_code == 201, resp.text
    body = resp.json()
    assert body["summary"]["docs"] == 6


def test_upload_folder_style(bundle: dict):
    assert bundle["summary"]["docs"] == 6
    assert bundle["summary"]["golden"] == 3


def test_zip_slip_rejected(client: TestClient, tmp_path: Path):
    zip_bytes = io.BytesIO()
    with zipfile.ZipFile(zip_bytes, "w") as zf:
        zf.writestr("../evil.txt", "pwned")
        zf.writestr("original/ok.png", b"\x89PNG\r\n")
    resp = client.post("/api/bundles", files={"files": ("evil.zip", zip_bytes.getvalue(), "application/zip")})
    assert resp.status_code == 201
    data_dir = tmp_path / "storage"
    assert not (data_dir / "evil.txt").exists()
    assert not list(data_dir.rglob("evil.txt"))


# ── missing/broken status ────────────────────────────────────────────────

def test_missing_and_broken(bundle: dict):
    docs = {d["id"]: d for d in bundle["docs"]}
    assert docs["PH003"]["has"]["harness"] is False
    assert docs["PH003"]["has"]["golden"] is False
    assert docs["OP004"]["has"]["preprocessed"] is False
    assert docs["DX006"]["has"] == {"original": True, "preprocessed": False, "ao_extract": False,
                                     "harness": False, "golden": False}
    assert docs["MC005"]["errors"], "broken harness json should be reported"
    assert "harness" in docs["MC005"]["errors"][0]


# ── compare statuses ─────────────────────────────────────────────────────

def test_compare_statuses(client: TestClient, bundle: dict):
    doc = client.get(f"/api/bundles/{bundle['id']}/docs/MC001").json()
    by_key = {(r["container"], r["row"], r["key"]): r for r in doc["compare"]}

    assert by_key[("", "", "발행일")]["ao_status"] == "MATCH"  # 2026-09-01 vs 20260901
    assert by_key[("", "", "공단부담총액")]["ao_status"] == "TYPE_MISMATCH"
    assert by_key[("", "", "공단부담총액")]["harness_status"] == "MATCH"

    row0 = by_key[("항목내역", 0, "금액")]
    assert row0["ao_status"] == "MISMATCH"  # 33000 != 35000
    assert row0["harness_status"] == "MATCH"  # harness 33000 -> 35000

    extra = [r for r in doc["compare"] if str(r["row"]).startswith("+")]
    assert extra, "extra table row (재료대) should be detected"
    assert any(r["ao_status"] == "EXTRA" for r in extra)
    assert any(r["harness_status"] == "EXTRA" for r in extra)

    doc2 = client.get(f"/api/bundles/{bundle['id']}/docs/DX002").json()
    by_key2 = {r["key"]: r for r in doc2["compare"]}
    assert by_key2["병명코드"]["ao_status"] == "MISMATCH"
    assert by_key2["병명코드"]["harness_status"] == "MATCH"
    assert by_key2["작성일자"]["ao_status"] == "MATCH"
    assert by_key2["작성일자"]["harness_status"] == "MISMATCH"  # 악화
    assert by_key2["병명"]["ao_status"] == "MISSING"
    assert by_key2["병명"]["harness_status"] == "MISSING"


def test_doc_detail_without_golden_matches_list_api(client: TestClient, bundle: dict):
    """Golden 없는 문서는 목록 API(bundle_view)처럼 compare=[], score=None 이어야 한다(가짜 EXTRA 불일치 금지)."""
    bid = bundle["id"]
    doc = client.get(f"/api/bundles/{bid}/docs/PH003").json()
    assert doc["has"]["golden"] is False
    assert doc["compare"] == []
    assert doc["score"] == {"ao": None, "harness": None}


def test_list_mismatch_matches_compare(client: TestClient, bundle: dict):
    """목록 API의 mismatch 는 상세 compare 에서 ao_status/harness_status 가 MATCH 가 아닌 행 수와 같아야 한다."""
    bid = bundle["id"]
    for summary in bundle["docs"]:
        if not summary["has"]["golden"]:
            continue
        doc = client.get(f"/api/bundles/{bid}/docs/{summary['id']}").json()
        expected = sum(
            1 for r in doc["compare"]
            if r["ao_status"] not in ("", "MATCH") or r["harness_status"] not in ("", "MATCH")
        )
        assert summary["mismatch"] == expected, summary["id"]


# ── golden create/save/delete ────────────────────────────────────────────

def test_create_golden_from_ao(client: TestClient, bundle: dict):
    bid = bundle["id"]
    resp = client.post(f"/api/bundles/{bid}/docs/PH003/golden", json={"from": "ao"})
    assert resp.status_code == 200, resp.text
    assert resp.json()["golden"]["documents"][0]["extracted_fields"]


def test_create_golden_409(client: TestClient, bundle: dict):
    bid = bundle["id"]
    resp = client.post(f"/api/bundles/{bid}/docs/MC001/golden", json={"from": "ao"})
    assert resp.status_code == 409


def test_create_golden_missing_source_404(client: TestClient, bundle: dict):
    bid = bundle["id"]
    resp = client.post(f"/api/bundles/{bid}/docs/PH003/golden", json={"from": "harness"})
    assert resp.status_code == 404


def test_golden_from_harness_strips_and_uses_final_value(client: TestClient, bundle: dict):
    bid = bundle["id"]
    d = client.delete(f"/api/bundles/{bid}/docs/MC001/golden")
    assert d.status_code == 204
    resp = client.post(f"/api/bundles/{bid}/docs/MC001/golden", json={"from": "harness"})
    assert resp.status_code == 200, resp.text
    golden = resp.json()["golden"]
    doc = golden["documents"][0]
    assert "harness" not in doc
    fields = {c["key"]: c for c in doc["extracted_fields"]}
    assert "harness" not in fields["공단부담총액"]
    assert fields["공단부담총액"]["value"] == "9000"  # final_value 로 대체됨
    table_cells = {c["key"]: c for c in doc["extracted_tables"][0]["rows"][0]}
    assert table_cells["금액"]["value"] == "35000"
    assert "harness" not in table_cells["금액"]


def test_put_golden_persists(client: TestClient, bundle: dict):
    bid = bundle["id"]
    new_golden = {"documents": [{"doc_type": "약제비영수증",
                                  "extracted_fields": [{"key": "조제일자", "value": "2026-09-05"}],
                                  "extracted_groups": [], "extracted_tables": []}]}
    r = client.put(f"/api/bundles/{bid}/docs/PH003/golden", json={"golden": new_golden})
    assert r.status_code == 200, r.text
    assert r.json()["golden"]["documents"][0]["extracted_fields"][0]["dtype"] == "string"  # 기본값 채움

    reloaded = client.get(f"/api/bundles/{bid}/docs/PH003").json()
    assert reloaded["golden"]["documents"][0]["extracted_fields"][0]["value"] == "2026-09-05"


def test_put_golden_invalid_422(client: TestClient, bundle: dict):
    bid = bundle["id"]
    r = client.put(f"/api/bundles/{bid}/docs/PH003/golden", json={"golden": {"nope": True}})
    assert r.status_code == 422


def test_delete_golden(client: TestClient, bundle: dict):
    bid = bundle["id"]
    r = client.delete(f"/api/bundles/{bid}/docs/OP004/golden")
    assert r.status_code == 204
    doc = client.get(f"/api/bundles/{bid}/docs/OP004").json()
    assert doc["has"]["golden"] is False


# ── AO/harness immutability ──────────────────────────────────────────────

def _hash_tree(root: Path) -> dict[str, str]:
    return {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(root.rglob("*")) if p.is_file()}


def test_ao_harness_files_never_modified(client: TestClient, bundle: dict, tmp_path: Path):
    bid = bundle["id"]
    bdir = tmp_path / "storage" / "bundles" / bid
    before = {**_hash_tree(bdir / "ao_extract"), **_hash_tree(bdir / "harness")}

    client.post(f"/api/bundles/{bid}/docs/PH003/golden", json={"from": "ao"})
    client.delete(f"/api/bundles/{bid}/docs/MC001/golden")
    client.post(f"/api/bundles/{bid}/docs/MC001/golden", json={"from": "harness"})
    client.put(f"/api/bundles/{bid}/docs/OP004/golden",
               json={"golden": {"documents": [{"doc_type": "x", "extracted_fields": [],
                                                "extracted_groups": [], "extracted_tables": []}]}})
    client.get(f"/api/bundles/{bid}/export/golden.xlsx")

    after = {**_hash_tree(bdir / "ao_extract"), **_hash_tree(bdir / "harness")}
    assert before == after


# ── review state ─────────────────────────────────────────────────────────

def test_review_state(client: TestClient, bundle: dict):
    bid = bundle["id"]
    r = client.put(f"/api/bundles/{bid}/docs/MC001/review", json={"review": "done"})
    assert r.status_code == 204
    doc = client.get(f"/api/bundles/{bid}/docs/MC001").json()
    assert doc["review"] == "done"

    view = client.get(f"/api/bundles/{bid}").json()
    assert view["summary"]["reviewed"] == 1

    client.put(f"/api/bundles/{bid}/docs/MC001/review", json={"review": ""})
    doc = client.get(f"/api/bundles/{bid}/docs/MC001").json()
    assert doc["review"] == ""


# ── export ───────────────────────────────────────────────────────────────

def test_export_zip(client: TestClient, bundle: dict):
    bid = bundle["id"]
    r = client.get(f"/api/bundles/{bid}/export/bundle.zip")
    assert r.status_code == 200
    with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
        names = zf.namelist()
    assert "golden/MC001.json" in names
    assert {n.split("/")[0] for n in names} >= {"original", "ao_extract", "harness", "golden"}

    r = client.get(f"/api/bundles/{bid}/export/bundle.zip", params={"doc": "MC001"})
    with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
        names = zf.namelist()
    assert names and all(n.split("/")[1].startswith("MC001.") for n in names)


def test_export_xlsx(client: TestClient, bundle: dict):
    bid = bundle["id"]
    r = client.get(f"/api/bundles/{bid}/export/golden.xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert set(wb.sheetnames) == {"요약", "필드", "표", "비교"}
    assert wb["요약"].max_row >= 2



def test_export_xlsx_without_growing_sheet_scans(client: TestClient, bundle: dict, monkeypatch):
    from openpyxl.worksheet.worksheet import Worksheet

    original = Worksheet.max_row.fget
    def max_row(ws):
        if ws.title in ("비교", "표"):
            raise AssertionError("export must not scan growing sheets for every row")
        return original(ws)

    with monkeypatch.context() as patch:
        patch.setattr(Worksheet, "max_row", property(max_row))
        response = client.get(f"/api/bundles/{bundle['id']}/export/golden.xlsx")
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content))
    for row in wb["표"]:
        if row[0].value:
            assert all(cell.font.bold for cell in row if cell.value is not None)
    from backend.export import _STATUS_FILL
    for row in wb["비교"].iter_rows(min_row=2):
        for index in (8, 10):
            cell = row[index]
            if cell.value in _STATUS_FILL:
                assert cell.fill == _STATUS_FILL[cell.value]


def test_enable_disable_docs(client: TestClient, bundle: dict):
    bid = bundle["id"]
    ids = [d["id"] for d in bundle["docs"]]
    assert all(d["enabled"] for d in bundle["docs"])
    assert bundle["summary_by_scope"]["disabled"]["docs"] == 0

    off = ids[:2]
    r = client.put(f"/api/bundles/{bid}/enabled", json={"ids": off, "enabled": False})
    assert r.status_code == 200
    view = r.json()
    assert [d["id"] for d in view["docs"] if not d["enabled"]] == off
    by = view["summary_by_scope"]
    assert by["disabled"]["docs"] == 2 and by["enabled"]["docs"] == len(ids) - 2
    assert view["summary"]["docs"] == len(ids)
    assert client.get(f"/api/bundles/{bid}/docs/{off[0]}").json()["enabled"] is False

    # 활성·비활성 ZIP 은 서로 겹치지 않는다
    def zip_docs(scope):
        r = client.get(f"/api/bundles/{bid}/export/bundle.zip", params={"scope": scope})
        assert r.status_code == 200
        with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
            return {Path(n).stem for n in zf.namelist()}
    assert zip_docs("disabled") == set(off)
    assert zip_docs("enabled") == set(ids) - set(off)
    r = client.get(f"/api/bundles/{bid}/export/bundle.zip", params={"scope": "nope"})
    assert r.status_code == 422

    def xlsx_docs(scope):
        r = client.get(f"/api/bundles/{bid}/export/golden.xlsx", params={"scope": scope})
        ws = load_workbook(io.BytesIO(r.content))["요약"]
        return {row[0] for row in ws.iter_rows(min_row=2, values_only=True)} - {"합계"}
    assert xlsx_docs("disabled") == set(off)
    assert xlsx_docs("enabled") == set(ids) - set(off)

    # 문서 단위 내보내기는 활성 여부와 무관
    r = client.get(f"/api/bundles/{bid}/export/bundle.zip", params={"doc": off[0]})
    with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
        assert zf.namelist()

    # 검수 상태 변경이 활성 여부를 지우지 않는다(같은 _state.json)
    client.put(f"/api/bundles/{bid}/docs/{off[0]}/review", json={"review": "done"})
    view = client.put(f"/api/bundles/{bid}/enabled", json={"ids": [off[1]], "enabled": True}).json()
    assert [d["id"] for d in view["docs"] if not d["enabled"]] == [off[0]]
    assert view["summary_by_scope"]["disabled"]["reviewed"] == 1

    assert client.put(f"/api/bundles/{bid}/enabled", json={"ids": ["없는문서"], "enabled": False}).status_code == 404
    assert client.put(f"/api/bundles/{bid}/enabled", json={"ids": off[0], "enabled": False}).status_code == 422
    assert client.put(f"/api/bundles/{bid}/enabled", json={"ids": off, "enabled": "no"}).status_code == 422


# ── images ───────────────────────────────────────────────────────────────

def test_image_endpoints(client: TestClient, bundle: dict):
    bid = bundle["id"]
    r = client.get(f"/api/bundles/{bid}/docs/MC001/image", params={"view": "original"})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("image/")

    r2 = client.get(f"/api/bundles/{bid}/docs/PH003/image", params={"view": "original", "page": 1})
    assert r2.status_code == 200
    assert r2.headers["x-pages"] == "2"

    r3 = client.get(f"/api/bundles/{bid}/docs/PH003/image", params={"view": "original", "page": 2})
    assert r3.status_code == 200
    assert r3.content != r2.content

    r4 = client.get(f"/api/bundles/{bid}/docs/DX006/image", params={"view": "preprocessed"})
    assert r4.status_code == 404


def test_raw_endpoint(client: TestClient, bundle: dict):
    bid = bundle["id"]
    r = client.get(f"/api/bundles/{bid}/docs/MC001/raw/ao_extract")
    assert r.status_code == 200
    assert json.loads(r.text)["transaction_id"] == "txn-mc001"

    r2 = client.get(f"/api/bundles/{bid}/docs/MC005/raw/harness")
    assert r2.status_code == 200  # 원문 그대로(파싱 실패여도)
    assert "not valid json" in r2.text

    r3 = client.get(f"/api/bundles/{bid}/docs/DX006/raw/golden")
    assert r3.status_code == 404


# ── delete bundle ────────────────────────────────────────────────────────

def test_delete_bundle(client: TestClient, bundle: dict):
    bid = bundle["id"]
    r = client.delete(f"/api/bundles/{bid}")
    assert r.status_code == 204
    assert client.get(f"/api/bundles/{bid}").status_code == 404
    assert bid not in [b["id"] for b in client.get("/api/bundles").json()]


def test_path_traversal_rejected(client: TestClient, bundle: dict):
    bid = bundle["id"]
    r = client.get(f"/api/bundles/{bid}/docs/..%2f..%2fetc/image")
    assert r.status_code in (400, 404, 422)


def test_unicode_doc_id(client: TestClient):
    """고객사 파일명(한글·공백·괄호·대괄호)에서 온 문서 ID도 상세·이미지를 연다."""
    doc_id = "[꾸미기]진단서 01 (1)"
    golden = json.dumps({"documents": [{"doc_type": "진단서", "extracted_fields": [], "extracted_groups": [], "extracted_tables": []}]})
    files = [("files", (f"b/original/{doc_id}.png", _png(), "image/png")),
             ("files", (f"b/golden/{doc_id}.answer.json", golden.encode(), "application/json"))]
    bid = client.post("/api/bundles", files=files).json()["id"]
    assert [d["id"] for d in client.get(f"/api/bundles/{bid}").json()["docs"]] == [doc_id]
    assert client.get(f"/api/bundles/{bid}/docs/{quote(doc_id)}").status_code == 200
    assert client.get(f"/api/bundles/{bid}/docs/{quote(doc_id)}/image").status_code == 200


def test_korean_zip_names(client: TestClient):
    """Windows 압축(CP949, UTF-8 플래그 없음)과 macOS NFD 파일명도 한글 폴더·문서 ID로 읽는다."""
    import unicodedata
    golden = json.dumps({"documents": []}).encode()
    items = [("원본/진단서 1.png", _png()), ("정답/진단서 1.json", golden)]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:  # ASCII 자리표시자로 쓴 뒤 CP949 바이트로 바꿔 UTF-8 플래그 없는 이름을 만든다
        for i, (name, data) in enumerate(items):
            zf.writestr(f"{i}".ljust(len(name.encode("cp949")), "_"), data)
        zf.writestr(unicodedata.normalize("NFD", "harness/진단서 1.harness.json"), golden)
    raw = buf.getvalue()
    for i, (name, _) in enumerate(items):
        raw = raw.replace(f"{i}".ljust(len(name.encode("cp949")), "_").encode(), name.encode("cp949"))
    buf = io.BytesIO(raw)
    bid = client.post("/api/bundles", files=[("files", ("b.zip", buf.getvalue(), "application/zip"))]).json()["id"]
    [doc] = client.get(f"/api/bundles/{bid}").json()["docs"]
    assert doc["id"] == "진단서 1"
    assert all(doc["has"][k] for k in ("original", "golden", "harness"))


def _png() -> bytes:
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (8, 8), "white").save(buf, format="PNG")
    return buf.getvalue()


def _harness_json(action: str, ao: str, title: str) -> bytes:
    doc = {"action": action, "ao_doc_type": ao, "title_doc_type": title, "title_line": "제목", "reason": "r"}
    return json.dumps({"documents": [], "harness": {"reclassification": {"documents": [doc]}}}).encode()


def test_doc_type_mismatch(client: TestClient):
    files = [("files", (f"harness/{i}.harness.json", _harness_json(a, "A", t), "application/json"))
             for i, a, t in (("RE1", "reextracted", "B"), ("KEPT", "kept", "B"), ("SAME", "reextracted", "A"))]
    bid = client.post("/api/bundles", files=files, data={"name": "mm"}).json()["id"]
    expected = {"ao": "A", "title": "B", "title_line": "제목", "reason": "r"}
    listed = {d["id"]: d["doc_type_mismatch"] for d in client.get(f"/api/bundles/{bid}").json()["docs"]}
    assert listed == {"RE1": expected, "KEPT": None, "SAME": None}
    assert client.get(f"/api/bundles/{bid}/docs/RE1").json()["doc_type_mismatch"] == expected
    assert client.get(f"/api/bundles/{bid}/docs/KEPT").json()["doc_type_mismatch"] is None


# ── 문서 종류 · 분류 채점 ─────────────────────────────────────────────────

def _typed(doc_type: str, value: str = "v") -> bytes:
    doc = {"doc_type": doc_type, "extracted_groups": [], "extracted_tables": [],
           "extracted_fields": [{"key": "발행일", "value": value, "dtype": "string"}]}
    return json.dumps({"documents": [doc]}).encode()


def _typed_bundle(client: TestClient) -> str:
    spec = {"OK": ("진료비영수증", "AC02922011", "진료비영수증"),   # 분류 정답
            "BAD": ("세부내역서", "AC02922011", "진료비세부산정내역서"),  # AO 오답, H 정답(별칭)
            "NOGOLD": (None, "AC02922011", "")}
    files = []
    for i, (g, a, h) in spec.items():
        for kind, t, sfx in (("golden", g, "answer"), ("ao_extract", a, "aiocr"), ("harness", h, "harness")):
            if t:
                files.append(("files", (f"{kind}/{i}.{sfx}.json", _typed(t, "x" if kind == "ao_extract" else "v"), "application/json")))
    return client.post("/api/bundles", files=files, data={"name": "cls"}).json()["id"]


def test_canon_and_apply_template():
    assert canon("Y000701333") == "소견서" and canon(" 입원확인서 ") == "입퇴원확인서"
    assert canon("모름") == canon("") == canon(None) == ""
    assert label("AC02922011") == label("진료비영수증") == "진료비영수증" and label("약제영수증") == label("Y000707300") == "약제비영수증"
    assert label("입원확인서") == "입퇴원확인서" and label("모름") == "모름" and label(None) == ""
    assert label("AC02922011", code=True) == "진료비영수증 (AC02922011)" and label("세부내역서", code=True) == "세부내역서"
    doc = {"extracted_fields": [{"key": "발행일", "value": "2024"}], "extracted_groups": [
        {"key": "환자정보", "fields": [{"key": "환자정보-성명", "value": "홍"}]}], "extracted_tables": [
        {"key": "항목내역", "headers": ["a"], "rows": [[{"key": "항목", "value": "주사"}, {"key": "zzz", "value": "1"}]]}]}
    out = apply_template(doc, "약제영수증")
    assert out["doc_type"] == "약제비영수증" and out is not TEMPLATES["약제비영수증"]
    out = apply_template(doc, "진료비영수증")
    cells = {c["key"]: c["value"] for c in out["extracted_fields"] + [f for g in out["extracted_groups"] for f in g["fields"]]}
    assert cells["발행일"] == "2024" and cells["환자정보-성명"] == "홍"
    rows = out["extracted_tables"][0]["rows"]
    assert [c["key"] for c in rows[0]] == ["항목"]


def test_create_golden_with_doc_type(client: TestClient):
    bid = _typed_bundle(client)
    r = client.post(f"/api/bundles/{bid}/docs/NOGOLD/golden", json={"from": "ao", "doc_type": "세부내역서"}).json()
    d0 = r["golden"]["documents"][0]
    assert d0["doc_type"] == "세부내역서"
    assert [c["key"] for c in d0["extracted_fields"]] == [c["key"] for c in TEMPLATES["세부내역서"]["extracted_fields"]]
    assert r["doc_type_suggest"] == "진료비영수증"
    assert r["doc_types"] == DOC_TYPES


def test_create_golden_same_type_keeps_ao(client: TestClient):
    bid = _typed_bundle(client)
    d0 = client.post(f"/api/bundles/{bid}/docs/NOGOLD/golden", json={"from": "ao", "doc_type": "진료비영수증"}).json()["golden"]["documents"][0]
    assert d0["doc_type"] == "진료비영수증" and d0["extracted_fields"][0]["value"] == "x"
    empty = client.post(f"/api/bundles/{bid}/docs/OK/golden", json={"from": "empty", "doc_type": "진단서"})
    assert empty.status_code == 409  # 이미 golden 있음


def test_classification_grading(client: TestClient):
    bid = _typed_bundle(client)
    docs = {d["id"]: d for d in client.get(f"/api/bundles/{bid}").json()["docs"]}
    assert docs["OK"]["classification"] == {"ao": True, "harness": True}
    assert docs["BAD"]["classification"] == {"ao": False, "harness": True}
    assert docs["NOGOLD"]["classification"] == {"ao": None, "harness": None}
    bad = client.get(f"/api/bundles/{bid}/docs/BAD").json()
    assert bad["classification"] == {"ao": False, "harness": True}
    assert bad["doc_type_by_source"] == {"golden": "세부내역서", "ao": "진료비영수증 (AC02922011)", "harness": "세부내역서"}
    # 목록·필터의 문서 유형은 표기(코드·이름·별칭)와 관계없이 표준 이름 하나로 묶인다
    assert {d["doc_type"] for d in docs.values()} == {"진료비영수증", "세부내역서"}
    assert docs["NOGOLD"]["doc_type"] == docs["OK"]["doc_type"] == bad["doc_type_by_source"]["ao"].split(" ")[0]
    summ = client.get(f"/api/bundles/{bid}").json()["summary"]
    assert summ["classification"]["ao"] == {"correct": 1, "total": 2, "accuracy": 0.5}
    assert summ["classification"]["harness"]["accuracy"] == 1.0
    # AO: BAD 는 필드 집계에서 제외 (OK 의 MISMATCH 1건만), H: 두 문서 모두 포함
    assert summ["score"]["ao"]["total"] == 1
    assert summ["score"]["harness"]["total"] == 2
    ws = load_workbook(io.BytesIO(client.get(f"/api/bundles/{bid}/export/golden.xlsx").content))["요약"]
    rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert rows["BAD"][2:4] == ("X", "O") and rows["OK"][2:4] == ("O", "O")
    assert rows["NOGOLD"][2:4] in (("", ""), (None, None))
    assert rows["합계"][5] == 0 and rows["합계"][6] == 1  # AO: BAD 제외 → OK 의 MATCH 0 / MISMATCH 1


def test_upload_hardening(client: TestClient):
    """전수조사에서 500을 내던 입력: 인코딩·형식이 어긋난 JSON, 잡파일, 손상 ZIP, 긴 파일명, 제어문자, '..' ID."""
    g = {"documents": [{"doc_type": "진단서", "extracted_fields": [{"key": "성명", "value": "홍\x0b길동"}],
                        "extracted_groups": [], "extracted_tables": []}]}
    files = [("files", ("b/original/a..b.png", _png(), "image/png")),
             ("files", ("b/golden/a..b.json", b"\xef\xbb\xbf" + json.dumps(g, ensure_ascii=False).encode(), "application/json")),
             ("files", ("b/harness/a..b.json", json.dumps({"documents": [{"doc_type": "진단서"}]}, ensure_ascii=False).encode("cp949"), "application/json")),
             ("files", ("b/ao_extract/a..b.json", b"[1, 2]", "application/json")),
             ("files", ("b/harness/Thumbs.db", b"\x00\xff\xfe", "application/octet-stream"))]
    bid = client.post("/api/bundles", files=files).json()["id"]
    [doc] = client.get(f"/api/bundles/{bid}").json()["docs"]
    assert doc["id"] == "a..b" and doc["has"]["golden"] and doc["has"]["harness"]
    assert [e.split(":")[0] for e in doc["errors"]] == ["ao_extract"]
    assert client.get("/api/bundles").status_code == 200
    assert client.get(f"/api/bundles/{bid}/docs/a..b").status_code == 200
    assert client.get(f"/api/bundles/{bid}/docs/a..b/raw/harness").json()["documents"][0]["doc_type"] == "진단서"
    assert client.get(f"/api/bundles/{bid}/export/golden.xlsx").status_code == 200
    assert client.put(f"/api/bundles/{bid}/docs/a..b/review", content=b"[1]").status_code == 422
    assert client.put(f"/api/bundles/{bid}/docs/a..b/golden", json={"golden": {"documents": [{"extracted_fields": [1]}]}}).status_code == 422

    assert client.post("/api/bundles", files=[("files", ("bad.zip", b"not a zip", "application/zip"))]).status_code == 400
    n = len(client.get("/api/bundles").json())
    assert client.post("/api/bundles", files=[("files", ("가" * 90 + ".png", _png(), "image/png"))]).status_code == 400
    assert len(client.get("/api/bundles").json()) == n  # 실패한 업로드는 번들을 남기지 않는다

    many = [("files", (f"m/original/{i}.png", b"x", "image/png")) for i in range(1100)]
    assert len(client.post("/api/bundles", files=many).json()["docs"]) == 1100


def test_upload_bad_zip_entry_leaves_no_bundle(client: TestClient):
    """암호·손상 ZIP 항목은 400이고 반쯤 만들어진 번들을 남기지 않는다."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("b/original/ok.png", _png())
        zf.writestr("b/original/bad.png", _png() * 50)
    corrupt = bytearray(buf.getvalue())
    corrupt[corrupt.index(b"bad.png") + len("bad.png") + 20] ^= 0xFF  # bad.png 압축 데이터 손상
    for blob in (bytes(corrupt), _encrypted_zip()):
        assert client.post("/api/bundles", files=[("files", ("x.zip", blob, "application/zip"))]).status_code == 400
        assert client.get("/api/bundles").json() == []


def _encrypted_zip() -> bytes:
    """암호 플래그(bit 0)를 켠 항목이 든 ZIP."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("b/original/enc.png", _png())
        zf.infolist()[0].flag_bits |= 0x1
    raw = bytearray(buf.getvalue())
    raw[raw.index(b"PK\x03\x04") + 6] |= 0x1  # 로컬 헤더 플래그도 맞춘다
    return bytes(raw)


def test_upload_content_length_over_limit_rejected_early(tmp_path: Path):
    client = TestClient(create_app(tmp_path / "storage", max_upload_mb=1))
    r = client.post("/api/bundles", files=[("files", ("b/original/a.png", b"x" * (1024 * 1024 + 1), "image/png"))])
    assert r.status_code == 413 and "upload too large" in r.json()["detail"]
    assert client.get("/api/bundles").json() == []


def test_upload_streams_without_loading_payload(client: TestClient, monkeypatch):
    """업로드 내용을 메모리로 읽지 않는다(UploadFile.read 미호출). 테스트 클라이언트가 본문을 들고 있어 tracemalloc은 쓰지 않는다."""
    import os
    from starlette.datastructures import UploadFile

    def boom(*a, **k):
        raise AssertionError("UploadFile.read called")
    monkeypatch.setattr(UploadFile, "read", boom)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_STORED) as zf:
        for i in range(10):
            zf.writestr(f"b/original/{i}.png", os.urandom(2 * 1024 * 1024))
    body = buf.getvalue()
    r = client.post("/api/bundles", files=[("files", ("b.zip", body, "application/zip"))])
    assert r.status_code == 201 and len(r.json()["docs"]) == 10


# ── 성능(캐시·썸네일·정적 버전) ────────────────────────────────────────────

def test_image_thumbnail_and_cache_header(client: TestClient, bundle: dict):
    url = f"/api/bundles/{bundle['id']}/docs/MC001/image"
    full = client.get(url)
    thumb = client.get(url, params={"w": 10})  # 64 로 clamp
    assert thumb.headers["content-type"] == "image/jpeg"
    assert thumb.headers["cache-control"] == full.headers["cache-control"] == "private, max-age=86400"
    from PIL import Image
    assert Image.open(io.BytesIO(thumb.content)).width <= 64
    assert client.get(url, params={"w": 10}).content == thumb.content  # 디스크 캐시 재사용


def test_caches_follow_file_changes(client: TestClient, bundle: dict):
    bid = bundle["id"]
    before = client.get(f"/api/bundles/{bid}").json()
    assert client.get(f"/api/bundles/{bid}").json() == before
    doc = next(d for d in before["docs"] if not d["has"]["golden"])["id"]
    assert client.post(f"/api/bundles/{bid}/docs/{doc}/golden", json={"from": "empty"}).status_code == 200
    after = client.get(f"/api/bundles/{bid}").json()
    assert next(d for d in after["docs"] if d["id"] == doc)["has"]["golden"]
    assert after["summary"]["golden"] == before["summary"]["golden"] + 1
    assert next(b for b in client.get("/api/bundles").json() if b["id"] == bid)["counts"]["golden"] == after["summary"]["golden"]


def test_static_version_and_gzip(client: TestClient):
    html = client.get("/").text
    ver = html.split('"/static/')[1].split("/")[0]
    assert client.get(f"/static/{ver}/app.js").headers["cache-control"].endswith("immutable")
    assert client.get("/static/app.js").headers["cache-control"] == "no-cache"
    assert client.get("/", headers={"accept-encoding": "gzip"}).headers["content-encoding"] == "gzip"


def test_golden_only_bundle_has_no_score_or_mismatch(client: TestClient, tmp_path: Path):
    """원본+Golden 만 올린 번들은 목록·상세 모두 점수 없음(None)·불일치 0 이어야 한다."""
    golden = {"documents": [{"doc_type": "세부내역서", "extracted_fields": [{"key": "사고발생일자", "value": "20230309", "dtype": "string"}],
                             "extracted_groups": [{"key": "합계", "fields": [{"key": "비급여총액", "value": None, "dtype": "string"}]}],
                             "extracted_tables": []}]}
    files = [("files", ("original/G001.png", b"\x89PNG\r\n\x1a\n", "image/png")),
             ("files", ("golden/G001.answer.json", json.dumps(golden).encode(), "application/json"))]
    resp = client.post("/api/bundles", files=files, data={"name": "golden-only"})
    assert resp.status_code == 201, resp.text
    summary = next(d for d in resp.json()["docs"] if d["id"] == "G001")
    assert summary["score"] == {"ao": None, "harness": None}
    assert summary["mismatch"] == 0
    doc = client.get(f"/api/bundles/{resp.json()['id']}/docs/G001").json()
    assert doc["score"] == {"ao": None, "harness": None}
    assert doc["compare"] and all(r["ao_status"] == "" and r["harness_status"] == "" for r in doc["compare"])


def test_export_xlsx_keeps_ocr_text_literal(client: TestClient, bundle: dict):
    bid = bundle["id"]
    did = bundle["docs"][0]["id"]
    golden = {"documents": [{"doc_type": "", "extracted_fields": [
        {"key": "=key\x01", "value": "=1+1\x02", "dtype": "string"}],
        "extracted_groups": [], "extracted_tables": []}]}
    assert client.put(f"/api/bundles/{bid}/docs/{did}/golden", json={"golden": golden}).status_code == 200
    response = client.get(f"/api/bundles/{bid}/export/golden.xlsx", params={"doc": did})
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content))
    assert wb["필드"]["E2"].value == "=key"
    assert wb["필드"]["F2"].value == "=1+1"
    assert wb["필드"]["E2"].data_type == wb["필드"]["F2"].data_type == "s"


@pytest.mark.parametrize("ext", ["bundle.zip", "golden.xlsx"])
def test_export_unknown_doc_returns_404(client: TestClient, bundle: dict, ext):
    response = client.get(f"/api/bundles/{bundle['id']}/export/{ext}", params={"doc": "unknown"})
    assert response.status_code == 404
